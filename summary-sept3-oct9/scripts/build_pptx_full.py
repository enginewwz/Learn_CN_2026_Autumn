#!/usr/bin/env python3
"""Full-text draft PPT: all report text and images, easy to cut down later."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parents[1]
FIG = BASE / "figures"
TOP = BASE.parent / "Sept17" / "report" / "figures"
OUT = BASE / "计网研讨课第一次实验总结-全量版.pptx"

INK = RGBColor(28, 37, 48)
MUTED = RGBColor(92, 105, 120)
LIGHT = RGBColor(248, 250, 252)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(220, 226, 234)
NAVY = RGBColor(15, 23, 42)
A = {
    1: RGBColor(47, 111, 176),
    2: RGBColor(46, 139, 87),
    3: RGBColor(201, 122, 30),
    4: RGBColor(111, 90, 168),
    5: RGBColor(185, 74, 72),
}
FONT = "Microsoft YaHei"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def font(run, size=10, color=INK, bold=False, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def text(slide, l, t, w, h, lines, size=9.5, color=INK, bold=False,
         align=PP_ALIGN.LEFT, spacing=1.04, after=2, name=FONT):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    if isinstance(lines, str):
        lines = [lines]
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        r = p.add_run()
        r.text = str(line)
        font(r, size=size, color=color, bold=bold, name=name)
    return box


def rect(slide, l, t, w, h, fill, line=None, lw=0.8):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(lw)
    return shape


def header(slide, kicker, title, accent, page):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    rect(slide, 0, 0, 13.333, 0.07, accent)
    text(slide, 0.55, 0.28, 9.5, 0.25, kicker, size=9.5, color=accent, bold=True)
    text(slide, 0.55, 0.52, 12.2, 0.55, title, size=22, color=INK, bold=True)
    rect(slide, 0.55, 6.99, 12.2, 0.01, LINE)
    text(slide, 0.55, 7.06, 9.5, 0.2,
         "计网研讨课第一次实验总结 · 吴峥 · 2024K80099013", size=8, color=MUTED)
    text(slide, 11.6, 7.06, 1.1, 0.2, str(page), size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def para_block(slide, l, t, w, h, heading, paras, accent, size=9.3, head_size=10.8):
    if heading:
        text(slide, l, t, w, 0.26, heading, size=head_size, color=accent, bold=True)
        t += 0.3
        h -= 0.3
    text(slide, l, t, w, h, paras, size=size, color=INK, spacing=1.03, after=2)


def caption(slide, l, t, w, s, size=8.2):
    text(slide, l, t, w, 0.35, s, size=size, color=MUTED, spacing=1.0)


def pic(slide, path, l, t, w, h, cap=None, cap_size=8.2):
    im = Image.open(path)
    iw, ih = im.size
    s = min(w / iw, h / ih)
    nw, nh = iw * s, ih * s
    x, y = l + (w - nw) / 2, t + (h - nh) / 2
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(nw), Inches(nh))
    if cap:
        caption(slide, l, t + h + 0.03, w, cap, cap_size)
    return y + nh


def table(slide, data, l, t, w, h, widths, accent, size=9.2, head_size=9.4):
    rows, cols = len(data), len(data[0])
    tbl = slide.shapes.add_table(rows, cols, Inches(l), Inches(t), Inches(w), Inches(h)).table
    tbl.horz_banding = False
    for i, cw in enumerate(widths):
        tbl.columns[i].width = Inches(cw)
    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = accent if r == 0 else (WHITE if r % 2 else RGBColor(247, 249, 252))
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            for j, line in enumerate(str(data[r][c]).split("\n")):
                if j:
                    p = tf.add_paragraph()
                p.space_after = Pt(0)
                r0 = p.add_run()
                r0.text = line
                font(r0, size=(head_size if r == 0 else size),
                     color=(WHITE if r == 0 else INK), bold=(r == 0))
    return tbl


def add_notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


# ============================================================
# 1 title
# ============================================================
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = NAVY
rect(s, 0.65, 1.1, 0.12, 2.3, A[1])
text(s, 0.95, 1.05, 10.8, 0.35, "计网研讨课 · 第一次实验总结", size=14, color=RGBColor(147, 197, 253), bold=True)
text(s, 0.95, 1.5, 11.5, 1.4, "从网络结构到网络性能", size=38, color=WHITE, bold=True)
text(s, 0.95, 2.8, 11.5, 0.5, "Wireshark · Socket · Switch · STP · Bufferbloat", size=16, color=RGBColor(203, 213, 225))
text(s, 0.95, 5.8, 5.5, 0.7, ["吴峥", "2024K80099013"], size=12, color=RGBColor(226, 232, 240), spacing=1.2)
text(s, 8.1, 5.8, 4.4, 0.5, "2026.09.03 — 2026.10.09", size=11, color=RGBColor(148, 163, 184), align=PP_ALIGN.RIGHT)

# ============================================================
# 2 overview
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "总览", "五次实验，实际上围绕两个主题", A[1], 2)
para_block(s, 0.55, 1.35, 6.0, 2.2, "两个主题",
           ["我们做这五次实验，实际上围绕两个主题：",
            "1. 理解网络的结构：分层、封装、协议栈，以及物理层 / 数据链路层 / 应用层设备与接口。",
            "2. 衡量并优化网络的性能：时延、吞吐、丢包、排队；Bufferbloat 与 AQM。"],
           A[1], size=10.5, head_size=12)
para_block(s, 6.85, 1.35, 5.9, 2.2, "一条主线",
           ["1、2 先看应用层和协议栈；3 进入数据链路层设备；4 解决交换机环路；5 看上层协议如何反过来影响下层队列。",
            "3 → 4 → 5 有递进：STP 解决环路；AQM 解决大缓存导致的高队列延迟。"],
           A[1], size=10.2, head_size=12)
data = [
    ["日期", "实验", "做了什么", "关键结果 / 关键词"],
    ["9/3", "互联网体系结构初识", "Mininet + Wireshark；2 万站点时延测量", "20,000 → 2,909；Total 58.6× c-latency"],
    ["9/10", "Socket 应用编程", "C + Linux Socket；HTTP/HTTPS 服务器；epoll 端口复用", "单线程事件循环 + 每连接状态机"],
    ["9/17", "Hub 与 Switch", "框架填空；广播 vs 自学习转发；环路实验", "广播风暴 423.4×；Switch 合计 2.12×"],
    ["9/24", "STP 生成树", "分布式竞选；端口状态；BPDU", "9 条链路 → 6 条树边，3 条阻塞"],
    ["10/9", "Bufferbloat", "maxq 扫描；taildrop / RED / CoDel；TSO/GSO", "RTT 652 → 57 ms；CoDel p95 59 ms"],
]
table(s, data, 0.55, 3.75, 12.2, 2.65, [0.7, 2.0, 5.0, 4.5], A[1], size=8.8, head_size=9.0)

# ============================================================
# 3 exp1 protocols
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 1 · 互联网体系结构初识", "一次网页访问里，能看到哪些协议？", A[1], 3)
data = [
    ["协议", "主要作用", "本次实验中的表现"],
    ["ARP", "根据同一链路中的 IP 地址查询对应的 MAC 地址", "出现 “who has” 请求和 “is at” 应答"],
    ["DNS", "将容易记忆的域名解析为 IP 地址", "www.ucas.ac.cn 解析为 124.16.77.5"],
    ["TCP", "在通信双方之间建立可靠的字节流连接", "与服务器 443 端口完成握手、传输和关闭"],
    ["HTTP", "规定浏览器或下载工具与 Web 服务器的请求和响应格式", "本次请求使用 HTTPS，HTTP 内容没有以明文显示"],
    ["HTTPS", "在 HTTP 与 TCP 之间加入 TLS 加密和身份验证", "Wireshark 中可以看到 TLS 握手及后续加密数据"],
]
table(s, data, 0.55, 1.35, 8.15, 4.75, [0.85, 3.6, 3.7], A[1], size=9.0)
para_block(s, 8.95, 1.35, 3.75, 3.0, "为什么先列这几个",
           ["它们是某条消息各自封装栈的最顶层。",
            "实际上可以继续向下展开：",
            "Ethernet < IP < UDP < DNS",
            "Ethernet < IP < TCP < TLS < HTTPS 数据",
            "这种展开直接对应网络的分层结构。"],
           A[1], size=8.9, head_size=10.2)
para_block(s, 8.95, 4.5, 3.75, 1.35, "ARP 的特殊之处",
           ["ARP 在实验 6 中也遇到：直接通过路由器发 ARP 请求。",
            "它直接作为以太网的载荷存在，没有更上层的封装。"],
           A[1], size=8.9, head_size=10.2)
para_block(s, 0.55, 6.2, 12.2, 0.6, "实验环境",
           ["WSL 中配置本机 mininet 环境；用 Wireshark 观察访问网站时产生的数据包，初步了解 ARP、DNS、TCP、HTTP、HTTPS 的作用，以及协议之间的分层与封装关系。"],
           A[1], size=8.7, head_size=9.8)

# ============================================================
# 4 exp1 timeline
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 1 · 一次访问的协议时间线", "TCP/TLS 既在顶层，也在别人的下层", A[1], 4)
data = [
    ["阶段", "包编号", "观察结果"],
    ["DNS 查询", "1～4", "h1 同时查询域名的 IPv4（A）和 IPv6（AAAA）地址，A 记录返回 124.16.77.5"],
    ["TCP 建立连接", "5～7", "客户端发送 SYN，服务器返回 SYN/ACK，客户端再发送 ACK，完成三次握手"],
    ["TLS 握手", "8～18", "双方协商加密参数并交换服务器证书，为 HTTPS 通信建立安全通道"],
    ["网页数据传输", "19～66", "服务器向客户端连续发送加密的网页数据，客户端发送 ACK 确认"],
    ["TCP 关闭连接", "67～69", "双方使用 FIN/ACK 结束本次连接"],
    ["ARP 交互", "70～75", "网络节点查询并应答局域网 IP 地址对应的 MAC 地址"],
]
table(s, data, 0.55, 1.35, 12.2, 4.0, [1.8, 1.1, 9.3], A[1], size=9.3)
para_block(s, 0.55, 5.5, 6.0, 1.4, "为什么 TCP/TLS 出现两次",
           ["握手阶段它们在封装栈顶层；传输阶段 HTTP 又建立在 TLS/TCP 之上。",
            "分层不是“一层只干一件事”，而是按阶段复用。"],
           A[1], size=9.2, head_size=10.5)
para_block(s, 6.75, 5.5, 6.0, 1.4, "一个可以提的细节",
           ["TCP 三次握手、四次挥手、TLS 握手不展开。",
            "但服务器把“确认客户端 FIN”和“发送自己的 FIN”合并在同一个 FIN, ACK 包中。"],
           A[1], size=9.2, head_size=10.5)

# ============================================================
# 5 exp1 method
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 1 · 时延测量", "真实互联网比“光速直线”慢多少？", A[1], 5)
para_block(s, 0.55, 1.35, 6.0, 3.1, "四个指标",
           ["Ping：向实际连接的服务器 IP 发送 3 次 ICMP Echo，取最小往返时间。",
            "DNS：完成域名解析所需的时间。",
            "TCP transfer：从收到 HTTP 响应的第一个字节，到接收完最后一个字节的时间。",
            "Total time：从开始请求到完整接收响应的总时间。"],
           A[1], size=10.2, head_size=12)
para_block(s, 0.55, 4.65, 6.0, 1.2, "实现流程",
           ["下数据库 → 读取网站列表 → 对每个网站执行访问和延迟测量 → 计算测量点到服务器的大圆距离 → 绘制图像。",
            "横轴 c-latency = 2 × distance / c，c = 299792.458 km/s，用真空光速估一个下界。"],
           A[1], size=9.2, head_size=10.8)
para_block(s, 6.85, 1.35, 5.9, 2.4, "数据量",
           ["根据抓到的网站列表访问网站，发生错误/超时的概率不小；小范围实验有效数据少，图毛刺多。",
            "最终租了一个阿里云服务器，花了约 3 小时，从 20,000 个网站中抓了 2,909 条有效数据。"],
           A[1], size=10.0, head_size=12)
para_block(s, 6.85, 3.95, 5.9, 2.9, "比较方式",
           ["实际耗时 ÷ c-latency。",
            "四条曲线：Ping / DNS / TCP transfer / Total HTTP time。",
            "中位数越接近 1×，越接近传播下界；越往右，说明协议栈和网络路径带来的额外开销越大。"],
           A[1], size=9.8, head_size=12)
para_block(s, 0.55, 6.0, 6.0, 0.8, "结果概览",
           ["真实互联网显著慢于真空光速；差距来自绕路、转发、DNS、TLS 和服务器处理。越往上层，额外开销越大。"],
           A[1], size=8.7, head_size=9.8)

# ============================================================
# 6 exp1 CDF
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 1 · 结果", "真实互联网显著慢于真空光速", A[1], 6)
pic(s, FIG / "01-latency-inflation-cdf.png", 0.55, 1.3, 7.4, 5.1,
    "四条曲线的中位数都在 1× 右侧；Ping 最靠左、最陡，Total 最靠右、长尾最长。")
para_block(s, 8.15, 1.3, 4.55, 2.4, "中位数",
           ["Ping：4.0×", "DNS：11.1×", "TCP transfer：4.5×", "Total HTTP：58.6×"],
           A[1], size=9.6, head_size=10.8)
para_block(s, 8.15, 3.35, 4.55, 3.2, "为什么",
           ["Ping 只走网络层/传输层，最接近传播下界，离散度也最小。",
            "DNS 受缓存和上游查询影响，离散大。",
            "TCP transfer 在目标服务器/缓存较近时衡量的是带宽，约 1/3 样本 < 1×。",
            "Total 叠加 DNS、建连、TLS、服务器处理和下载，所以倍数最大。"],
           A[1], size=8.8, head_size=10.5)

# ============================================================
# 7 exp2 intro
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 2 · Socket 应用编程", "用 socket 把 HTTP/HTTPS 请求当成文件来处理", A[2], 7)
# topology from the report
text(s, 0.55, 1.25, 4.0, 0.25, "实验拓扑", size=10, color=A[2], bold=True)
pic(s, FIG / "sept10-topology.png", 0.55, 1.62, 5.9, 1.35,
    "拓扑：h2 → s1 → h1；s1 是 OVS 交换机，h1-s1 链路 100 Mbit/s、10 ms。")
para_block(s, 0.55, 3.35, 5.9, 2.35, "socket 是什么 / 实验两步",
           ["socket 接口为应用层提供了一个传输层（TCP/UDP）的类文件系统抽象。",
            "1. 框架已经给定了 socket 生命周期，允许我们像处理普通文件一样处理每个请求；按状态码表格逐项处理 HTTP/HTTPS 文本。",
            "2. 实验要求多请求并行处理；这里选择 epoll 端口复用，而不是多线程。",
            "挑战：对网络库不熟时，理解 epoll 并不轻松；它和直接多线程的方案区别很大。"],
           A[2], size=9.4, head_size=10.5)
para_block(s, 6.75, 1.25, 6.0, 3.2, "epoll 的难点",
           ["不是“多开几个线程”那么简单：一个线程里同时维护多个连接。",
            "每个连接有自己的状态、读缓冲、写缓冲和文件进度。",
            "事件就绪时处理；做不下去就保存现场、改关注事件、返回。",
            "端口复用更像“把任务切碎再合并”：自己维护进入/返回点，调度交给 epoll。"],
           A[2], size=9.6, head_size=10.8)
para_block(s, 6.75, 4.55, 6.0, 1.6, "和 pthread 的关系",
           ["如果局限于单核来理解，和 pthread 处理起来类似；不同是你需要自己维护一个内部的读写缓冲，处理清楚进入和返回的过程。",
            "调度其实也是上层维护好的（主要指睡眠和唤醒的过程）。"],
           A[2], size=8.8, head_size=10.2)
para_block(s, 0.55, 5.85, 5.9, 0.75, "本实验实现内容",
           ["HTTP/HTTPS 文件服务器；200/301/400/404/405/416/431；单段 Range；TLS。"],
           A[2], size=8.4, head_size=9.6)

# ============================================================
# 8 exp2 struct
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 2 · 连接对象", "一个连接，就是一个显式状态机", A[2], 8)
pic(s, FIG / "connection-code.png", 0.55, 1.3, 5.3, 3.8)
data = [
    ["字段", "作用"],
    ["fd、tls、ssl", "保存套接字、是否使用 TLS 以及对应的 OpenSSL 连接对象"],
    ["state", "记录当前处于握手、读取、发送或关闭阶段；监听对象使用 LISTENER"],
    ["input、used", "保存已收到的请求数据及有效长度"],
    ["output、output_len、\nsent", "保存当前待发送的数据块、块长度和已发送长度"],
    ["file_fd、file_offset、\nfile_remaining", "保存打开的文件、下一次读取位置及尚未读入缓冲区的字节数"],
]
table(s, data, 6.15, 1.3, 6.5, 3.85, [1.7, 4.8], A[2], size=8.4)
para_block(s, 0.55, 5.3, 5.3, 1.45, "为什么放在一起",
           ["状态、缓冲、文件进度都跟着同一个对象走；释放时统一处理，不会漏。"],
           A[2], size=9.2, head_size=10.2)
para_block(s, 6.15, 5.3, 6.5, 1.45, "归属",
           ["每个 client 一份 connection；两个端口不共享数据缓冲。",
            "监听对象也有 connection（state = LISTENER），accept 后才新建真正的连接。"],
           A[2], size=9.2, head_size=10.2)

# ============================================================
# 9 exp2 loop
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 2 · epoll 事件循环", "等待事件 → 循环处理 → 分派 listener / connection", A[2], 9)
pic(s, FIG / "02-socket-epoll-loop.png", 0.55, 1.25, 12.2, 3.6,
    "epoll_wait 阻塞等待，一次最多 64 个就绪事件；处理完一个事件回到循环头，一组完成才回等待。")
para_block(s, 0.55, 5.25, 4.0, 1.55, "listener 分支",
           ["接受新连接 → 新建 connection 结构体并注册到 epoll → drive。"],
           A[2], size=9.0, head_size=10.2)
para_block(s, 4.75, 5.25, 4.0, 1.55, "connection 分支",
           ["直接 drive；从上次停下的状态继续。"],
           A[2], size=9.0, head_size=10.2)
para_block(s, 8.95, 5.25, 3.8, 1.55, "做不下去",
           ["记住进度、改 epoll 关注事件、返回等待。",
            "一次最多 64 个事件；每个监听口每轮最多接受 64 个；listen backlog 128。"],
           A[2], size=8.6, head_size=10.2)

# ============================================================
# 10 exp2 FSM
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 2 · 每连接状态机", "drive 一次可以连续走多个状态，不是每次都回等待", A[2], 10)
pic(s, FIG / "02-socket-fsm-automaton.png", 0.45, 1.25, 9.5, 5.0)
para_block(s, 10.05, 1.25, 2.75, 2.1, "状态",
           ["HANDSHAKE → READING → WRITING → CLOSING → RELEASE"],
           A[2], size=8.4, head_size=9.5)
para_block(s, 10.05, 3.35, 2.75, 2.9, "等待 / EAGAIN",
           ["EAGAIN / WANT_* 不是错误：现在做不了。",
            "保存现场，改 epoll 关注事件，下次 drive 继续。",
            "READING 收齐请求头后可以直接进 WRITING；只有做不下去或完成才返回等待。"],
           A[2], size=8.2, head_size=9.5)
para_block(s, 0.55, 6.35, 12.2, 0.55, None,
           ["socket 允许一个 fd 又读又写：read() 从接收缓冲区头部取已到达字节，write() 追加到发送缓冲区尾部；内核维护各自游标。"],
           A[2], size=8.5)

# ============================================================
# 11 exp3 framework
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 3 · Hub 与 Switch", "先读懂框架：侵入式链表和“填空”方法", A[3], 11)
pic(s, FIG / "list-code.png", 0.55, 1.25, 7.4, 3.9)
para_block(s, 8.2, 1.25, 4.55, 2.1, "侵入式链表",
           ["list_head 嵌在对象里，不需要额外节点；list_entry 用成员地址反推整个对象。"],
           A[3], size=9.2, head_size=10.5)
para_block(s, 8.2, 3.5, 4.55, 1.65, "读框架的方法",
           ["先找 TODO 和调用链，不必通读全部代码；go to definition + 全局搜索。"],
           A[3], size=9.2, head_size=10.5)
para_block(s, 0.55, 5.35, 12.2, 1.4, "难点 / 注意",
           ["实验 3 的复杂度主要来自框架，而不是转发逻辑；结构体和宏已经预定义好，但刚上手不容易理解。",
            "代码细节本身比较容易，主要是有锁编程需要注意。"],
           A[3], size=9.4, head_size=10.8)

# ============================================================
# 12 exp3 forwarding
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 3 · 交换机怎么转发", "Hub 广播；Switch 学 MAC、查表、单播", A[3], 12)
para_block(s, 0.55, 1.3, 5.8, 2.1, "Hub（物理层）",
           ["没有 MAC 地址；收到帧就向其他所有端口复制。",
            "集线器运行在物理层，因此相对上层交换机的传输是透明的。"],
           A[3], size=9.5, head_size=11)
para_block(s, 0.55, 3.55, 5.8, 2.9, "Switch（数据链路层）",
           ["每个接口有一个唯一的 MAC 地址（区别于一般的每个交换机共享一个）；整机一张全局 MAC 地址表。",
            "收到帧：查目的 MAC，命中则只从对应端口发出；未命中则向其他端口泛洪；同时学习源 MAC -> 入接口。",
            "表项动态插入，30 s 老化；拓扑变化时重新学习。"],
           A[3], size=9.0, head_size=11)
para_block(s, 6.75, 1.3, 6.0, 2.4, "一个容易混淆的点",
           ["iface_info_t 里每个接口各有一份 MAC；",
            "mac_port_map 是整机一张学习表，不是“交换机 MAC”。",
            "每个表项：源 MAC + iface* + 最近访问时间。"],
           A[3], size=9.2, head_size=10.8)
para_block(s, 6.75, 3.95, 6.0, 2.5, "为什么需要 STP",
           ["交换机通过 MAC 地址确定向哪个端口转发；但遇到环路时，广播帧会不断复制。",
            "没有 STP，交换机也无法避免泛洪成环。",
            "集线器只适合放在网络最末端。"],
           A[3], size=9.2, head_size=10.8)

# ============================================================
# 13 exp3 storm
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 3 · 环路", "一次 ping 如何在 6 秒内变成 20 万个帧", A[3], 13)
pic(s, TOP / "topology_ring.png", 0.55, 1.3, 3.55, 2.2, "环形拓扑：三个 Hub 首尾相连。")
pic(s, FIG / "03-ring-broadcast-storm.png", 4.35, 1.3, 5.1, 3.2,
    "包速率约 33 kpps、吞吐贴着 20 Mbit/s，直到抓包结束都没有衰减。")
pic(s, FIG / "03-ring-frame-composition.png", 9.65, 1.3, 3.0, 2.0,
    "ARP 41.2%；ICMP Echo Reply 58.5%；Echo Request 只有 0.2%。")
para_block(s, 0.55, 5.0, 8.8, 1.8, "为什么停不下来",
           ["一次 ping 在 6 秒内被放大成 20 万个帧，放大 423.4×；唯一签名的 Echo Request 只有 1 个，却复制出 483 个副本。",
            "Hub 不学 MAC，广播帧在环里来回复制；帧没有 TTL，不会自己消失。STP 从逻辑上断开环路。"],
           A[3], size=9.0, head_size=10.8)
para_block(s, 9.65, 4.0, 3.0, 2.6, "数字",
           ["423.4× 放大", "33 kpps", "20 Mbit/s", "41.2% ARP", "58.5% ICMP Echo Reply"],
           A[3], size=9.0, head_size=10.5)

# ============================================================
# 14 exp3 throughput
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 3 · Hub vs Switch", "广播的代价：多条流共享同一条下行链路", A[3], 14)
pic(s, TOP / "topology_three_nodes.png", 0.55, 1.3, 5.0, 3.0, "简单拓扑：两台主机经 Hub / Switch 通信。")
pic(s, FIG / "03-hub-vs-switch-throughput.png", 5.75, 1.3, 7.0, 4.4,
    "场景 A：Hub 两条流 5.85 / 2.99 Mbit/s；Switch 各自接近 9.4 Mbit/s。")
para_block(s, 0.55, 5.0, 12.2, 1.8, "结论",
           ["交换机能够做到点对点通信，而集线器因为运行在物理层只能广播；Scenario A 合计吞吐约 2.12×。",
            "广播导致的链路吞吐瓶颈比较直观；两个实验也说明集线器只适合运行在整个网络最末端。",
            "环形拓扑下如果没有合适手段规避，会发生的泛洪；这直接引出了实验 4。"],
           A[3], size=9.2, head_size=10.8)

# ============================================================
# 15 exp4 STP
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 4 · STP 生成树", "每个端口保存自己的最优结果，主机仲裁根端口", A[4], 15)
pic(s, FIG / "04-seven-node-topology.png", 0.55, 1.3, 5.4, 4.7,
    "七节点拓扑：最终只保留一棵无环生成树。")
para_block(s, 6.25, 1.3, 6.4, 2.0, "分布式竞选",
           ["每台交换机先认为自己是根；交换 BPDU 后逐步收敛。",
            "交换机的每个端口储存各自获得的最优结果；如果收到的包比全局存储的更优，由主机仲裁重选根端口。"],
           A[4], size=8.9, head_size=10.5)
para_block(s, 6.25, 3.45, 6.4, 1.8, "比较顺序",
           ["根桥 ID -> 到根开销 -> 指定桥 ID -> 指定端口 ID。",
            "最后两项保证任意两个向量都能分出唯一胜负。"],
           A[4], size=8.9, head_size=10.5)
para_block(s, 6.25, 5.35, 6.4, 1.45, "一个细节 / 结果",
           ["路径开销只加一次：在交换机层加，避免两个端口对同一段链路重复加代价。",
            "9 条链路 -> 6 条树边，3 条阻塞。"],
           A[4], size=8.7, head_size=10.5)

# ============================================================
# 16 exp5 bufferbloat
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 5 · Bufferbloat", "缓冲越大，RTT 越高，但吞吐并没有提高", A[5], 16)
pic(s, FIG / "05-topology.png", 0.55, 1.3, 4.2, 3.5, "10 Mbit 瓶颈 + maxq 缓冲；下半部分是 AQM 对比拓扑。")
pic(s, FIG / "05-bufferbloat-sweep.png", 5.0, 1.3, 7.7, 4.8,
    "maxq = 10 / 50 / 200 / 1000：缓冲越大，cwnd / 队列 / RTT 锯齿越大、周期越长。")
para_block(s, 0.55, 5.0, 4.2, 1.8, "数字",
           ["maxq 1000 -> 50：RTT 中位 652 -> 57 ms，吞吐仍是 9.5 Mbit/s。",
            "再压到 10 包（小于 BDP）才损失约 15% 吞吐。"],
           A[5], size=8.8, head_size=10.5)
para_block(s, 5.0, 6.05, 7.7, 0.8, "解释",
           ["TCP AIMD：拥塞窗口/队列/RTT 呈线性上升、断崖下降；这是上层协议对下层设备性能的跨层耦合。"],
           A[5], size=8.6, head_size=10.2)

# ============================================================
# 17 exp5 AQM
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 5 · AQM", "延迟可以独立于吞吐来优化", A[5], 17)
pic(s, FIG / "05-aqm-rtt.png", 0.55, 1.3, 6.0, 3.7,
    "RTT：taildrop 在低速相位冲到 5–8 s；RED 约 0.8 s；CoDel 10–50 ms。")
pic(s, FIG / "05-aqm-throughput.png", 6.75, 1.3, 5.9, 3.7,
    "吞吐：虚线是服务端线上速率；三种规则几乎完全重合。")
para_block(s, 0.55, 5.25, 12.2, 1.55, "结论",
           ["taildrop p95 6.1 s，RED 0.8 s，CoDel 59 ms；三者总交付量都约 298 MB。",
            "CoDel 用“排队停留时间”而不是“队列长度”做控制：以包在队列中的停留时间为控制量，比单纯把队列排满效率高。"],
           A[5], size=9.3, head_size=10.8)

# ============================================================
# 18 exp5 TSO
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "实验 5 · 一个意外", "TSO/GSO 会把 mininet 队列里的包“放大”", A[5], 18)
pic(s, FIG / "05-tso-comparison.png", 0.55, 1.3, 6.3, 4.2,
    "taildrop：TSO/GSO 开 vs 关，队列字节和 RTT 相差一个数量级。")
para_block(s, 7.2, 1.3, 5.5, 2.0, "为什么",
           ["开启 TSO/GSO 时，网络包由物理网卡切片；在 mininet 模拟中，相当于队列里每个包都发生膨胀，因而造成更严重的 bufferbloat。"],
           A[5], size=9.3, head_size=10.8)
para_block(s, 7.2, 3.5, 5.5, 2.2, "为什么 RED / CoDel 不受太大影响",
           ["RED 的阈值 avpkt 是字节量；CoDel 的控制量是排队时间。",
            "它们不按“包个数”做判断。"],
           A[5], size=9.3, head_size=10.8)
para_block(s, 0.55, 5.8, 12.2, 0.9, None,
           ["这也是为什么同一个 bufferbloat 实验，在开启/关闭 offload 时曲线会差一个数量级。"],
           A[5], size=9.0)

# ============================================================
# 19 summary
# ============================================================
s = prs.slides.add_slide(BLANK)
header(s, "总结", "五次实验放在一张表里", A[1], 19)
data = [
    ["#", "实验", "实验目的"],
    ["1", "互联网体系结构初识", "理解网络分层结构 + 了解衡量网络性能的指标"],
    ["2", "Socket 编程", "理解应用层如何处理 HTTP/HTTPS 数据 + 优化端口复用效率"],
    ["3", "Hub 与 Switch", "理解物理层和数据链路层设备和功能"],
    ["4", "STP 生成树", "理解数据链路层设备功能"],
    ["5", "Bufferbloat", "理解上层协议对下层设备的影响，优化数据链路层设备性能"],
]
table(s, data, 0.55, 1.3, 12.2, 3.3, [0.55, 3.0, 8.65], A[1], size=9.5)
para_block(s, 0.55, 4.95, 6.0, 1.7, "3 → 4 → 5",
           ["STP 解决交换机的环路问题；AQM 解决拥塞协议导致的大缓存下交换机高队列延迟问题。"],
           A[3], size=9.6, head_size=10.8)
para_block(s, 6.75, 4.95, 6.0, 1.7, "两个收获",
           ["始终在明确的某个网络层级上学习，分别加深对网络结构和性能的理解。",
            "分层结构对抽象有优势，但分层不能解耦性能问题。"],
           A[1], size=9.3, head_size=10.8)

# ============================================================
# 20 closing
# ============================================================
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = NAVY
rect(s, 0.65, 1.3, 0.12, 1.7, A[2])
text(s, 0.95, 1.25, 10.8, 0.4, "最后", size=12, color=RGBColor(147, 197, 253), bold=True)
text(s, 0.95, 1.7, 11.4, 1.0, "结构上看清层级，性能上看到跨层耦合", size=28, color=WHITE, bold=True)
text(s, 0.95, 3.05, 11.4, 1.2,
     ["每个实验都落在明确的网络层级上；", "但 TCP 的 AIMD、队列和缓存又会把不同层绑在一起。"],
     size=14, color=RGBColor(203, 213, 225), spacing=1.2)
text(s, 0.95, 6.25, 11.4, 0.6, "谢谢", size=20, color=WHITE, bold=True)
text(s, 8.2, 6.38, 4.3, 0.4, "吴峥 · 2024K80099013", size=9.5, color=RGBColor(148, 163, 184), align=PP_ALIGN.RIGHT)

prs.save(str(OUT))
print("saved", OUT)
