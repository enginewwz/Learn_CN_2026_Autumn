#!/usr/bin/env python3
"""Build a basic editable PPTX from report-txt.md content."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parents[1]
FIG = BASE / "figures"
TOP = BASE.parent / "Sept17" / "report" / "figures"
OUT = BASE / "计网研讨课第一次实验总结.pptx"

# palette
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(100, 116, 139)
BG = RGBColor(249, 250, 252)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(226, 232, 240)
NAVY = RGBColor(15, 23, 42)
ACCENT = {
    1: RGBColor(47, 111, 176),
    2: RGBColor(46, 139, 87),
    3: RGBColor(201, 122, 30),
    4: RGBColor(111, 90, 168),
    5: RGBColor(185, 74, 72),
}
TINT = {
    1: RGBColor(235, 243, 251),
    2: RGBColor(235, 247, 241),
    3: RGBColor(252, 244, 232),
    4: RGBColor(243, 239, 250),
    5: RGBColor(251, 238, 238),
}
FONT = "Microsoft YaHei"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def set_font(run, size=14, color=INK, bold=False, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def add_text(slide, left, top, width, height, lines, size=14, color=INK,
             bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             line_spacing=1.08, space_after=3, name=FONT, margins=(0.02, 0.02)):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(margins[0])
    tf.margin_right = Inches(margins[0])
    tf.margin_top = Inches(margins[1])
    tf.margin_bottom = Inches(margins[1])
    if isinstance(lines, str):
        lines = [lines]
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = str(line)
        set_font(r, size=size, color=color, bold=bold, name=name)
    return box


def add_rect(slide, left, top, width, height, fill, line=None, radius=False, lw=1.0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(lw)
    if radius:
        try:
            shape.adjustments[0] = 0.08
        except Exception:
            pass
    shape.shadow.inherit = False
    return shape


def add_card(slide, left, top, width, height, title, body, accent, tint,
             title_size=13, body_size=10.5, body_color=MUTED):
    add_rect(slide, left, top, width, height, tint, accent, radius=True, lw=1.0)
    add_text(slide, left + 0.16, top + 0.12, width - 0.3, 0.34,
             title, size=title_size, color=INK, bold=True)
    if body:
        add_text(slide, left + 0.16, top + 0.52, width - 0.3, height - 0.6,
                 body, size=body_size, color=body_color, line_spacing=1.12,
                 space_after=2)


def add_kicker_title(slide, kicker, title, accent, page=None):
    # background
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    # top accent line
    add_rect(slide, 0, 0, 13.333, 0.08, accent)
    add_text(slide, 0.65, 0.34, 9.0, 0.3, kicker, size=10.5, color=accent, bold=True)
    add_text(slide, 0.65, 0.63, 11.9, 0.72, title, size=26, color=INK, bold=True)
    # footer
    add_rect(slide, 0.65, 7.03, 12.02, 0.012, LINE)
    add_text(slide, 0.65, 7.1, 8.5, 0.22,
             "计网研讨课第一次实验总结 · 吴峥 · 2024K80099013",
             size=8.5, color=MUTED)
    if page is not None:
        add_text(slide, 11.7, 7.1, 1.0, 0.22, str(page), size=8.5,
                 color=MUTED, align=PP_ALIGN.RIGHT)


def add_table(slide, data, left, top, width, height, col_widths, accent,
              font_size=10.5, header_size=10.5):
    rows, cols = len(data), len(data[0])
    table = slide.shapes.add_table(rows, cols, Inches(left), Inches(top),
                                   Inches(width), Inches(height)).table
    table.horz_banding = False
    table.vert_banding = False
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.045)
            cell.margin_bottom = Inches(0.045)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = accent
            else:
                cell.fill.fore_color.rgb = WHITE if r % 2 == 1 else RGBColor(247, 249, 252)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(0)
            chunks = str(data[r][c]).split("\n")
            for j, chunk in enumerate(chunks):
                if j:
                    p = tf.add_paragraph()
                    p.alignment = PP_ALIGN.LEFT
                    p.space_after = Pt(0)
                run = p.add_run()
                run.text = chunk
                set_font(run, size=(header_size if r == 0 else font_size),
                         color=(WHITE if r == 0 else INK), bold=(r == 0))
    return table


def fit_image(slide, path, left, top, max_w, max_h, border=False):
    img = Image.open(path)
    iw, ih = img.size
    scale = min(max_w / iw, max_h / ih)
    w, h = iw * scale, ih * scale
    x = left + (max_w - w) / 2
    y = top + (max_h - h) / 2
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                   width=Inches(w), height=Inches(h))
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(0.75)
    return x, y, w, h


def add_caption(slide, left, top, width, text, size=8.8):
    add_text(slide, left, top, width, 0.3, text, size=size, color=MUTED,
             line_spacing=1.0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_stat(slide, left, top, width, height, value, label, accent):
    add_rect(slide, left, top, width, height, WHITE, LINE, radius=True, lw=0.8)
    add_text(slide, left + 0.12, top + 0.12, width - 0.24, 0.42,
             value, size=20, color=accent, bold=True)
    add_text(slide, left + 0.12, top + 0.56, width - 0.24, height - 0.6,
             label, size=8.8, color=MUTED, line_spacing=1.05)


# ---------------------------------------------------------------------------
# 1. title
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
add_rect(slide, 0.65, 1.15, 0.12, 2.2, ACCENT[1])
add_text(slide, 0.95, 1.12, 10.6, 0.4, "计网研讨课 · 第一次实验总结",
         size=15, color=RGBColor(147, 197, 253), bold=True)
add_text(slide, 0.95, 1.6, 11.3, 1.4, "从网络结构到网络性能",
         size=40, color=WHITE, bold=True)
add_text(slide, 0.95, 2.82, 11.3, 0.55,
         "Wireshark · Socket · Switch · STP · Bufferbloat",
         size=17, color=RGBColor(203, 213, 225))
add_text(slide, 0.95, 5.75, 5.8, 0.8,
         ["吴峥", "2024K80099013"],
         size=13, color=RGBColor(226, 232, 240), line_spacing=1.2)
add_text(slide, 8.2, 5.75, 4.4, 0.8,
         ["2026.09.03 — 2026.10.09"], size=12,
         color=RGBColor(148, 163, 184), align=PP_ALIGN.RIGHT)
# five small color blocks
for i, (label, color) in enumerate([
    ("01 体系结构", ACCENT[1]), ("02 Socket", ACCENT[2]), ("03 Switch", ACCENT[3]),
    ("04 STP", ACCENT[4]), ("05 Bufferbloat", ACCENT[5]),
]):
    x = 0.95 + i * 2.34
    add_rect(slide, x, 6.65, 2.04, 0.06, color)
    add_text(slide, x, 6.78, 2.04, 0.3, label, size=8.5,
             color=RGBColor(148, 163, 184))

# ---------------------------------------------------------------------------
# 2. overview
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "总览", "五次实验，其实只围绕两个问题", ACCENT[1], 2)
add_card(slide, 0.65, 1.55, 5.75, 1.75, "理解网络的结构",
         ["分层、封装、协议栈", "物理层 / 数据链路层 / 应用层设备与接口"], ACCENT[1], TINT[1],
         title_size=15, body_size=11)
add_card(slide, 6.9, 1.55, 5.75, 1.75, "衡量与优化网络的性能",
         ["时延、吞吐、丢包、排队", "Bufferbloat 与 AQM：延迟能否独立于吞吐优化？"], ACCENT[5], TINT[5],
         title_size=15, body_size=11)

items = [
    ("9/3", "互联网体系结构初识", "Wireshark + 2 万站点时延测量", ACCENT[1], TINT[1]),
    ("9/10", "Socket 应用编程", "HTTP/HTTPS 服务器 + epoll 端口复用", ACCENT[2], TINT[2]),
    ("9/17", "Hub 与 Switch", "广播 vs 自学习转发；环路广播风暴", ACCENT[3], TINT[3]),
    ("9/24", "STP 生成树", "分布式竞选出一棵无环树", ACCENT[4], TINT[4]),
    ("10/9", "Bufferbloat", "队列、AQM、TSO/GSO 的影响", ACCENT[5], TINT[5]),
]
for i, (date, title, desc, color, tint) in enumerate(items):
    x = 0.65 + i * 2.42
    add_rect(slide, x, 3.65, 2.24, 2.45, tint, color, radius=True, lw=1.0)
    add_text(slide, x + 0.14, 3.8, 1.95, 0.3, date, size=10, color=color, bold=True)
    add_text(slide, x + 0.14, 4.18, 1.95, 0.9, title, size=13.5, color=INK, bold=True,
             line_spacing=1.0)
    add_text(slide, x + 0.14, 5.15, 1.95, 0.8, desc, size=8.7, color=MUTED,
             line_spacing=1.1)
add_text(slide, 0.65, 6.45, 12.0, 0.35,
         "3 → 4 → 5 有递进：STP 解决环路；AQM 解决大缓存导致的高队列延迟。",
         size=10.5, color=MUTED)

# ---------------------------------------------------------------------------
# 3. exp1 protocols
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 1 · 互联网体系结构初识", "一次网页访问里，能看到哪些协议？", ACCENT[1], 3)
add_text(slide, 0.65, 1.45, 6.4, 0.3,
         "Mininet + Wireshark：先抓包，再按封装关系拆协议。",
         size=11, color=MUTED)
data = [
    ["协议", "主要作用", "本次实验中的表现"],
    ["ARP", "同一链路内：IP → MAC", "出现 “who has” 请求 / “is at” 应答"],
    ["DNS", "域名 → IP 地址", "www.ucas.ac.cn → 124.16.77.5"],
    ["TCP", "可靠的字节流连接", "与服务器 443 端口完成握手、传输、关闭"],
    ["HTTP", "浏览器与 Web 服务器的请求/响应格式", "本次走 HTTPS，HTTP 明文不可见"],
    ["HTTPS", "在 HTTP 与 TCP 之间加入 TLS 加密和身份验证", "TLS 握手 + 后续加密数据"],
]
add_table(slide, data, 0.65, 1.85, 7.7, 4.6, [0.95, 3.0, 3.75], ACCENT[1], font_size=10.0)
add_card(slide, 8.65, 1.85, 4.0, 4.6, "它们可以继续向下展开",
         ["Ethernet < IP < UDP < DNS",
          "Ethernet < IP < TCP < TLS < HTTPS 数据",
          "",
          "ARP 是个例外：它直接作为以太网载荷，没有更上层的封装。",
          "",
          "能展开，正说明协议是分层的。"],
         ACCENT[1], TINT[1], title_size=12.5, body_size=10.2)
add_notes(slide, "Wireshark 里先看协议列，再逐层展开 Ethernet / IP / TCP / TLS。ARP 没有 IP 层之上的封装。")

# ---------------------------------------------------------------------------
# 4. exp1 timeline
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 1 · 一次访问的协议时间线", "TCP/TLS 既在顶层，也在别人的下层", ACCENT[1], 4)
data = [
    ["阶段", "包编号", "观察结果"],
    ["DNS 查询", "1–4", "同时查 A / AAAA；A 记录返回 124.16.77.5"],
    ["TCP 建立连接", "5–7", "SYN → SYN/ACK → ACK，三次握手"],
    ["TLS 握手", "8–18", "协商加密参数、交换服务器证书"],
    ["网页数据传输", "19–66", "服务器连续发送加密数据；客户端 ACK"],
    ["TCP 关闭连接", "67–69", "FIN/ACK 结束；服务器把 ACK 与自己的 FIN 合并"],
    ["ARP 交互", "70–75", "查询并应答局域网 IP 对应的 MAC"],
]
add_table(slide, data, 0.65, 1.65, 12.0, 4.35, [2.1, 1.2, 8.7], ACCENT[1], font_size=10.5)
add_card(slide, 0.65, 6.2, 12.0, 0.62, "为什么 TCP/TLS 出现在两个位置？",
         "握手阶段它们在协议栈顶层；传输阶段 HTTP 又建立在 TLS/TCP 之上。分层不是固定的“一层只干一件事”，而是按阶段复用。",
         ACCENT[1], TINT[1], title_size=10.5, body_size=9.4)
add_notes(slide, "FIN/ACK 合并是一个可以提的细节：服务器确认客户端 FIN 时，把自己的 FIN 也带上了。")

# ---------------------------------------------------------------------------
# 5. exp1 method
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 1 · 时延测量", "真实互联网比“光速直线”慢多少？", ACCENT[1], 5)
metrics = [
    ("Ping", "向实际服务器 IP 发 3 次 ICMP Echo，取最小 RTT"),
    ("DNS", "完成域名解析所需的时间"),
    ("TCP transfer", "从 HTTP 响应第一个字节，到接收完最后一个字节"),
    ("Total time", "从开始请求，到完整接收响应的总时间"),
]
for i, (name, desc) in enumerate(metrics):
    y = 1.65 + i * 1.02
    add_rect(slide, 0.65, y, 6.9, 0.86, WHITE, LINE, radius=True, lw=0.8)
    add_text(slide, 0.85, y + 0.1, 1.75, 0.3, name, size=11.5, color=ACCENT[1], bold=True)
    add_text(slide, 2.55, y + 0.1, 4.8, 0.62, desc, size=9.6, color=MUTED, line_spacing=1.08)
add_card(slide, 7.85, 1.65, 4.8, 1.55, "横轴：c-latency",
         "c-latency = 2 × distance / c\nc = 299792.458 km/s\n用大圆距离估一个“真空光速下界”。",
         ACCENT[1], TINT[1], title_size=12.5, body_size=10.0)
add_card(slide, 7.85, 3.45, 4.8, 1.6, "数据量",
         "20,000 个网站 → 2,909 条有效数据\n租阿里云服务器，约 3 小时抓取\n小范围实验毛刺多，所以扩大样本。",
         ACCENT[1], TINT[1], title_size=12.5, body_size=10.0)
add_card(slide, 7.85, 5.25, 4.8, 1.35, "比较方式",
         "实际耗时 ÷ c-latency。越接近 1×，越接近传播下界。",
         ACCENT[1], TINT[1], title_size=12.5, body_size=10.0)

# ---------------------------------------------------------------------------
# 6. exp1 CDF
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 1 · 结果", "真实互联网显著慢于真空光速", ACCENT[1], 6)
x, y, w, h = fit_image(slide, FIG / "01-latency-inflation-cdf.png", 0.65, 1.5, 8.4, 4.7)
stats = [("4.0×", "Ping 中位数", ACCENT[1]), ("11.1×", "DNS 中位数", ACCENT[2]),
         ("4.5×", "TCP transfer 中位数", ACCENT[3]), ("58.6×", "Total HTTP 中位数", ACCENT[5])]
for i, (v, label, color) in enumerate(stats):
    add_stat(slide, 9.25, 1.55 + i * 1.18, 3.4, 1.0, v, label, color)
add_caption(slide, 0.65, 6.25, 8.4,
            "四条曲线的中位数都在 1× 右侧；Ping 最靠左、最陡，Total 最靠右、长尾最长。")
add_text(slide, 9.25, 6.25, 3.4, 0.6,
         "TCP transfer 约 1/3 样本 < 1×，被截断在左边界。",
         size=8.5, color=MUTED)
add_notes(slide, "Ping 最接近传播下界；DNS 受缓存影响离散大；TCP 在近处衡量带宽；Total 叠加 DNS、建连、TLS、服务器处理和下载。")

# ---------------------------------------------------------------------------
# 7. exp2 intro
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 2 · Socket 应用编程", "用 socket 把 HTTP/HTTPS 请求当成文件来处理", ACCENT[2], 7)
add_text(slide, 0.65, 1.55, 4.9, 0.3, "实验拓扑", size=11, color=ACCENT[2], bold=True)
fit_image(slide, FIG / "sept10-topology.png", 0.65, 1.72, 5.9, 1.32)
add_caption(slide, 0.65, 3.08, 5.9,
            "拓扑：h2 → s1 → h1；s1 是 OVS 交换机，h1-s1 链路 100 Mbit/s、10 ms。")
add_card(slide, 0.65, 3.55, 5.9, 1.35, "socket 是什么",
         "应用层看到的传输层抽象：像文件一样 read / write，底层由内核维护收发缓冲。",
         ACCENT[2], TINT[2], title_size=12.5, body_size=10.2)
add_card(slide, 0.65, 5.1, 5.9, 1.35, "实验分两步",
         "1. 框架已给定 socket 生命周期；按状态码处理 HTTP/HTTPS 文本。\n2. 多请求并行：选择 epoll 端口复用，而不是 pthread。",
         ACCENT[2], TINT[2], title_size=12.5, body_size=10.0)
add_card(slide, 6.85, 1.6, 5.8, 2.45, "真正的挑战：epoll",
         ["不是“多开几个线程”那么简单。",
          "• 一个线程里同时维护多个连接",
          "• 每个连接有自己的状态、读/写缓冲、文件进度",
          "• 事件就绪时处理；做不下去就保存现场、改关注事件、返回"],
         ACCENT[2], TINT[2], title_size=13.0, body_size=10.0)
add_card(slide, 6.85, 4.2, 5.8, 2.25, "epoll 三件事",
         ["注册：告诉内核关注哪些 fd 的哪些事件",
          "等待：阻塞到有事件就绪（一次最多 64 个）",
          "处理：逐个处理；没完成就改关注事件"],
         ACCENT[2], TINT[2], title_size=13.0, body_size=10.0)
add_notes(slide, "socket 提供类文件系统抽象；epoll 与 pthread 的差异主要在状态和缓冲要自己维护。")

# ---------------------------------------------------------------------------
# 8. exp2 connection struct
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 2 · 连接对象", "一个连接，就是一个显式状态机", ACCENT[2], 8)
fit_image(slide, FIG / "connection-code.png", 0.65, 1.65, 5.7, 3.55)
add_card(slide, 0.65, 5.4, 5.7, 1.1, "为什么放在一起",
         "状态、缓冲、文件进度都在同一个对象里；释放时统一处理，不会漏。",
         ACCENT[2], TINT[2], title_size=10.5, body_size=9.2)
data = [
    ["字段", "作用"],
    ["fd / tls / ssl", "套接字、是否 TLS、OpenSSL 连接对象"],
    ["state", "握手 / 读取 / 发送 / 关闭；监听对象是 LISTENER"],
    ["input / used", "已收到的请求数据与有效长度"],
    ["output / output_len / sent", "当前待发数据、块长度、已发长度"],
    ["file_fd / offset / remaining", "打开的文件、下一次读取位置、剩余字节"],
]
add_table(slide, data, 6.75, 1.65, 5.9, 3.55, [1.75, 4.15], ACCENT[2], font_size=9.2)
add_card(slide, 6.75, 5.4, 5.9, 1.1, "归属",
         "每个 client 一份；状态、缓冲、文件进度都跟着对象走，不是两个端口共享。",
         ACCENT[2], TINT[2], title_size=10.5, body_size=9.4)

# ---------------------------------------------------------------------------
# 9. exp2 epoll loop
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 2 · epoll 事件循环", "等待事件 → 循环处理 → 分派给 listener / connection", ACCENT[2], 9)
fit_image(slide, FIG / "02-socket-epoll-loop.png", 0.65, 1.45, 12.0, 3.65)
add_caption(slide, 0.65, 5.2, 12.0,
            "epoll_wait 阻塞等待，一次最多 64 个就绪事件；处理完一个事件回到循环头，一组完成才回等待。")
add_card(slide, 0.65, 5.6, 3.85, 1.1, "listener 分支",
         "接受新连接 → 新建 connection 结构体并注册 → drive",
         ACCENT[2], TINT[2], title_size=10.5, body_size=9.5)
add_card(slide, 4.72, 5.6, 3.85, 1.1, "connection 分支",
         "直接 drive；从上次停下的状态继续",
         ACCENT[2], TINT[2], title_size=10.5, body_size=9.5)
add_card(slide, 8.79, 5.6, 3.85, 1.1, "做不下去",
         "记住进度、改 epoll 关注事件、返回等待",
         ACCENT[2], TINT[2], title_size=10.5, body_size=9.5)
add_notes(slide, "事件循环与连接状态机分开看：这里讲的是“什么时候处理哪个连接”。")

# ---------------------------------------------------------------------------
# 10. exp2 FSM
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 2 · 每连接状态机", "drive 一次可以连续走多个状态，不是每次都回等待", ACCENT[2], 10)
fit_image(slide, FIG / "02-socket-fsm-automaton.png", 0.65, 1.35, 12.0, 5.25)
add_text(slide, 0.65, 6.62, 5.9, 0.35,
         "不是每次转移都回等待：READING 收齐请求头 → 直接 WRITING。",
         size=9.8, color=MUTED)
add_text(slide, 6.75, 6.62, 5.9, 0.35,
         "EAGAIN / WANT_*：保存现场，改 epoll 关注事件，下次 drive 继续。",
         size=9.8, color=MUTED)

# ---------------------------------------------------------------------------
# 11. exp3 framework
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 3 · Hub 与 Switch", "先读懂框架：侵入式链表和“填空”方法", ACCENT[3], 11)
fit_image(slide, FIG / "list-code.png", 0.65, 1.65, 7.1, 3.5)
add_card(slide, 8.05, 1.65, 4.6, 1.55, "侵入式链表",
         "list_head 嵌在对象里，不需要额外节点；list_entry 用成员地址反推对象。",
         ACCENT[3], TINT[3], title_size=11.5, body_size=9.3)
add_card(slide, 8.05, 3.35, 4.6, 1.55, "读框架的方法",
         "先找 TODO 和调用链，不必通读全部代码；go to definition + 全局搜索。",
         ACCENT[3], TINT[3], title_size=11.5, body_size=9.3)
add_card(slide, 0.65, 5.35, 12.0, 1.15, "这一步的难点",
         "结构体和宏已经预定义好，但刚上手不容易理解；实验 3 的复杂度主要来自框架，而不是转发逻辑本身。",
         ACCENT[3], TINT[3], title_size=11.5, body_size=9.8)

# ---------------------------------------------------------------------------
# 12. exp3 forwarding
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 3 · 交换机怎么转发", "Hub 广播；Switch 学 MAC、查表、单播", ACCENT[3], 12)
add_card(slide, 0.65, 1.55, 5.9, 1.3, "Hub（物理层）",
         "没有 MAC 地址；收到帧就向其他所有端口复制。",
         ACCENT[3], TINT[3], title_size=12.5, body_size=10.2)
add_card(slide, 0.65, 3.0, 5.9, 1.3, "Switch（数据链路层）",
         "每个接口有自己的 MAC；整机一张全局 MAC 地址表。",
         ACCENT[3], TINT[3], title_size=12.5, body_size=10.2)
add_card(slide, 6.75, 1.55, 5.9, 2.55, "收到一帧之后",
         ["1. 查目的 MAC：命中 → 只从对应端口发出",
          "2. 未命中 → 向其他端口泛洪",
          "3. 学习源 MAC → 入接口，写入 MAC 表",
          "4. 表项动态插入，30 s 老化",
          "交换机对上层透明：上层只看 IP，不关心端口。"],
         ACCENT[3], TINT[3], title_size=13.0, body_size=9.9)
add_card(slide, 6.75, 4.3, 5.9, 1.7, "MAC 表项",
         ["每个表项：源 MAC + iface* + 最近访问时间",
          "256 个哈希桶；30 s 老化；拓扑变化后重新学习。"],
         ACCENT[3], TINT[3], title_size=12.0, body_size=9.3)
add_card(slide, 0.65, 4.55, 5.9, 1.55, "一个容易混淆的点",
         "iface_info_t 里每个接口各有一份 MAC；\nmac_port_map 是整机一张学习表，不是“交换机 MAC”。",
         ACCENT[3], TINT[3], title_size=11.0, body_size=9.3)

# ---------------------------------------------------------------------------
# 13. exp3 ring storm
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 3 · 环路", "一次 ping 如何在 6 秒内变成 20 万个帧", ACCENT[3], 13)
fit_image(slide, TOP / "topology_ring.png", 0.65, 1.5, 3.55, 2.25)
add_caption(slide, 0.65, 3.8, 3.55, "环形拓扑：三个 Hub 首尾相连。")
fit_image(slide, FIG / "03-ring-broadcast-storm.png", 4.45, 1.45, 5.0, 3.15)
add_caption(slide, 4.45, 4.65, 5.0, "包速率约 33 kpps、吞吐贴着 20 Mbit/s，直到抓包结束都没有衰减。")
fit_image(slide, FIG / "03-ring-frame-composition.png", 9.7, 1.45, 3.0, 2.0)
add_caption(slide, 9.7, 3.5, 3.0, "ARP 41.2%；ICMP Echo Reply 58.5%；真正的 Echo Request 只有 0.2%。")
add_stat(slide, 9.7, 4.0, 3.0, 1.05, "423.4×", "6 秒内帧数放大倍数", ACCENT[3])
add_text(slide, 9.7, 5.25, 3.0, 1.45,
         "唯一签名的 Echo Request 只有 1 个，却复制出 483 个副本。\n\n没有 STP，交换机也无法避免泛洪成环。",
         size=9.2, color=MUTED, line_spacing=1.12)
add_card(slide, 0.65, 5.55, 8.7, 1.2, "为什么停不下来",
         "Hub 不学 MAC，广播帧在环里来回复制；帧没有 TTL，不会自己消失。STP 才能从逻辑上断开环路。",
         ACCENT[3], TINT[3], title_size=11.0, body_size=9.3)

# ---------------------------------------------------------------------------
# 14. exp3 hub vs switch
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 3 · Hub vs Switch", "广播的代价：多条流共享同一条下行链路", ACCENT[3], 14)
fit_image(slide, TOP / "topology_three_nodes.png", 0.65, 1.55, 5.0, 3.0)
add_caption(slide, 0.65, 4.65, 5.0, "简单拓扑：两台主机经 Hub / Switch 通信。")
fit_image(slide, FIG / "03-hub-vs-switch-throughput.png", 5.9, 1.5, 6.8, 4.5)
add_caption(slide, 5.9, 6.05, 6.8, "场景 A：Hub 两条流 5.85 / 2.99 Mbit/s；Switch 各自接近 9.4 Mbit/s。")
add_card(slide, 0.65, 5.05, 5.0, 1.65, "结论",
         "Switch 点对点转发，合计吞吐约 2.12×；\nHub 把帧复制到所有端口，适合放在网络最末端。",
         ACCENT[3], TINT[3], title_size=11.5, body_size=9.6)

# ---------------------------------------------------------------------------
# 15. exp4 STP
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 4 · STP 生成树", "每个端口保存自己的最优结果，主机仲裁根端口", ACCENT[4], 15)
fit_image(slide, FIG / "04-seven-node-topology.png", 0.65, 1.55, 5.6, 4.7)
add_caption(slide, 0.65, 6.3, 5.6, "七节点拓扑：最终只保留一棵无环生成树。")
add_card(slide, 6.55, 1.55, 6.1, 1.55, "分布式竞选",
         "每台交换机先认为自己是根；交换 BPDU 后，按优先级逐步收敛。",
         ACCENT[4], TINT[4], title_size=11.5, body_size=9.8)
add_card(slide, 6.55, 3.25, 6.1, 1.85, "比较顺序",
         "根桥 ID → 到根开销 → 指定桥 ID → 指定端口 ID。\n最后两项保证任意两个向量都能分出唯一胜负。",
         ACCENT[4], TINT[4], title_size=11.5, body_size=9.8)
add_card(slide, 6.55, 5.25, 6.1, 1.2, "一个细节",
         "路径开销只能加一次：在交换机层加，避免两个端口对同一段链路重复加代价。",
         ACCENT[4], TINT[4], title_size=11.5, body_size=9.3)

# ---------------------------------------------------------------------------
# 16. exp5 bufferbloat
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 5 · Bufferbloat", "缓冲越大，RTT 越高，但吞吐并没有提高", ACCENT[5], 16)
fit_image(slide, FIG / "05-topology.png", 0.65, 1.5, 4.3, 3.6)
add_caption(slide, 0.65, 5.2, 4.3, "10 Mbit 瓶颈 + maxq 缓冲；下半部分是 AQM 对比拓扑。")
fit_image(slide, FIG / "05-bufferbloat-sweep.png", 5.2, 1.45, 7.45, 4.75)
add_caption(slide, 5.2, 6.25, 7.45, "maxq = 10 / 50 / 200 / 1000：缓冲越大，cwnd / 队列 / RTT 锯齿越大、周期越长。")
add_stat(slide, 0.65, 5.55, 2.05, 1.1, "652→57 ms", "maxq 1000→50，RTT 中位", ACCENT[5])
add_stat(slide, 2.9, 5.55, 2.05, 1.1, "9.5 Mbit/s", "吞吐不变", ACCENT[5])

# ---------------------------------------------------------------------------
# 17. exp5 AQM
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 5 · AQM", "延迟可以独立于吞吐来优化", ACCENT[5], 17)
fit_image(slide, FIG / "05-aqm-rtt.png", 0.65, 1.55, 6.0, 3.7)
add_caption(slide, 0.65, 5.35, 6.0, "RTT：taildrop 在低速相位冲到 5–8 s；RED 约 0.8 s；CoDel 10–50 ms。")
fit_image(slide, FIG / "05-aqm-throughput.png", 6.85, 1.55, 5.8, 3.7)
add_caption(slide, 6.85, 5.35, 5.8, "吞吐：虚线是服务端线上速率；三种规则几乎完全重合。")
add_card(slide, 0.65, 5.8, 12.0, 0.9, "结论",
         "taildrop p95 6.1 s，RED 0.8 s，CoDel 59 ms；三者总交付量都约 298 MB。CoDel 用“排队停留时间”而不是“队列长度”做控制。",
         ACCENT[5], TINT[5], title_size=11.0, body_size=9.3)

# ---------------------------------------------------------------------------
# 18. exp5 TSO
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "实验 5 · 一个意外", "TSO/GSO 会把 mininet 队列里的包“放大”", ACCENT[5], 18)
fit_image(slide, FIG / "05-tso-comparison.png", 0.65, 1.55, 6.6, 4.25)
add_caption(slide, 0.65, 5.9, 6.6, "taildrop：TSO/GSO 开 vs 关，队列字节和 RTT 相差一个数量级。")
add_card(slide, 7.55, 1.6, 5.1, 2.0, "为什么",
         "开启 TSO/GSO 时，网络包由物理网卡切片；在 mininet 的软件队列里，相当于每个包都膨胀。",
         ACCENT[5], TINT[5], title_size=11.5, body_size=9.8)
add_card(slide, 7.55, 3.8, 5.1, 2.0, "为什么 RED / CoDel 不受太大影响",
         "RED 的 avpkt 阈值按字节；CoDel 的控制量是排队时间。\n它们不按“包个数”做判断。",
         ACCENT[5], TINT[5], title_size=11.5, body_size=9.8)

# ---------------------------------------------------------------------------
# 19. summary table
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
add_kicker_title(slide, "总结", "五次实验放在一张表里", ACCENT[1], 19)
data = [
    ["#", "实验", "实验目的"],
    ["1", "互联网体系结构初识", "理解网络分层结构；了解衡量网络性能的指标"],
    ["2", "Socket 编程", "理解应用层如何处理 HTTP/HTTPS；优化端口复用效率"],
    ["3", "Hub 与 Switch", "理解物理层和数据链路层设备与功能"],
    ["4", "STP 生成树", "理解数据链路层设备如何消除环路"],
    ["5", "Bufferbloat", "理解上层协议对下层设备的影响；优化队列延迟"],
]
add_table(slide, data, 0.65, 1.65, 12.0, 3.55, [0.55, 3.2, 8.25], ACCENT[1], font_size=10.5)
add_card(slide, 0.65, 5.45, 5.8, 1.3, "3 → 4 → 5 的递进",
         "STP 解决交换机的环路；AQM 解决大缓存导致的高队列延迟。",
         ACCENT[3], TINT[3], title_size=11.5, body_size=9.7)
add_card(slide, 6.85, 5.45, 5.8, 1.3, "两个收获",
         "分层抽象降低理解成本；但分层不能解耦性能问题。",
         ACCENT[1], TINT[1], title_size=11.5, body_size=9.7)

# ---------------------------------------------------------------------------
# 20. closing
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
add_rect(slide, 0.65, 1.35, 0.12, 1.7, ACCENT[2])
add_text(slide, 0.95, 1.3, 10.8, 0.45, "最后", size=13, color=RGBColor(147, 197, 253), bold=True)
add_text(slide, 0.95, 1.75, 11.4, 1.0, "结构上看清层级，性能上看到跨层耦合",
         size=30, color=WHITE, bold=True)
add_text(slide, 0.95, 3.1, 11.4, 1.2,
         ["每个实验都落在明确的网络层级上；",
          "但 TCP 的 AIMD、队列和缓存又会把不同层绑在一起。"],
         size=15, color=RGBColor(203, 213, 225), line_spacing=1.25)
add_text(slide, 0.95, 6.35, 11.4, 0.6, "谢谢", size=22, color=WHITE, bold=True)
add_text(slide, 8.2, 6.45, 4.3, 0.4, "吴峥 · 2024K80099013", size=10,
         color=RGBColor(148, 163, 184), align=PP_ALIGN.RIGHT)

prs.save(str(OUT))
print("saved", OUT)
